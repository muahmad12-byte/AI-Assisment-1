"""
AI Resume ATS Analyzer
----------------------
Upload a resume (PDF / DOCX / TXT) and get:
  * an overall ATS score (0-100) with a category breakdown
  * strengths, problems found, and concrete improvements
  * missing keywords, and (optionally) a match score against a job description

UI: Streamlit    |    AI: Google Gemini Flash (google-genai SDK)
"""

import hashlib
import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# ----------------------------------------------------------------------------
# Config
# ----------------------------------------------------------------------------
# "gemini-flash-latest" is an alias that always points to Google's current Flash model,
# so the app keeps working when older models are retired.
DEFAULT_MODEL = os.getenv("GEMINI_MODEL", "gemini-flash-latest")
# Tried in order if the chosen model is not found (404 / retired).
FALLBACK_MODELS = ["gemini-flash-latest", "gemini-3.7-flash", "gemini-3.5-flash",
                   "gemini-3.1-flash-lite"]
MAX_FILE_MB = 5
MAX_RESUME_CHARS = 20000  # keeps the prompt small and fast
MIN_RESUME_CHARS = 150  # below this the file is probably scanned / empty

CATEGORIES = {
    "formatting": "Formatting & ATS-friendliness",
    "keywords": "Keywords & Skills",
    "impact": "Impact & Achievements",
    "structure": "Structure & Sections",
    "readability": "Readability & Language",
}

SYSTEM_PROMPT = """You are an expert technical recruiter and ATS (Applicant Tracking System) \
specialist. You evaluate resumes strictly and honestly. Do not inflate scores: an average \
resume should land between 50 and 70.

Score each category from 0 to 100:
- formatting: clean single-column layout, standard headings, no tables/graphics clues, \
consistent dates, contact info present
- keywords: relevant hard/soft skills, industry terms, tools and technologies
- impact: quantified achievements (numbers, %, $), strong action verbs, results over duties
- structure: standard sections (Summary, Experience, Education, Skills, Projects), logical order, \
appropriate length
- readability: concise bullets, grammar, no typos, consistent tense, no filler

If a job description is provided, also compute "job_match_score" (0-100) and list \
"missing_keywords" taken from the job description that are absent from the resume. \
If no job description is provided, set "job_match_score" to null and give \
"missing_keywords" as important keywords commonly expected for the candidate's apparent role.

Respond with ONLY valid JSON (no markdown fences, no commentary) using exactly this schema:
{
  "overall_score": <int 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "detected_role": "<best guess of the target role>",
  "category_scores": {
    "formatting": <int>, "keywords": <int>, "impact": <int>,
    "structure": <int>, "readability": <int>
  },
  "job_match_score": <int or null>,
  "strengths": ["..."],
  "issues": [{"severity": "high|medium|low", "section": "...", "problem": "..."}],
  "improvements": [{"priority": "high|medium|low", "suggestion": "...", "example": "..."}],
  "missing_keywords": ["..."],
  "rewritten_bullets": [{"original": "...", "improved": "..."}]
}
Give 3-6 strengths, 4-8 issues, 5-8 improvements and up to 4 rewritten bullets. \
"example" must show a concrete before/after or sample text. Only use facts present in the \
resume; never invent employers, degrees or numbers in rewritten bullets - use placeholders \
like [X%] where a metric is needed."""


# ----------------------------------------------------------------------------
# File parsing
# ----------------------------------------------------------------------------
def extract_text(file_name: str, data: bytes) -> str:
    """Extract plain text from PDF, DOCX or TXT bytes."""
    ext = os.path.splitext(file_name)[1].lower()

    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password-protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        text = "\n".join(pages)

    elif ext == ".docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        # many resumes keep content inside tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        text = "\n".join(parts)

    elif ext == ".txt":
        text = data.decode("utf-8", errors="ignore")

    else:
        raise ValueError(f"Unsupported file type: {ext or 'unknown'}")

    return clean_text(text)


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ----------------------------------------------------------------------------
# Gemini
# ----------------------------------------------------------------------------
def parse_json_response(raw: str) -> dict:
    """Parse model output into a dict, tolerating markdown fences / extra text."""
    if not raw:
        raise ValueError("The model returned an empty response.")
    cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw.strip(), flags=re.IGNORECASE)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        start, end = cleaned.find("{"), cleaned.rfind("}")
        if start != -1 and end > start:
            return json.loads(cleaned[start : end + 1])
        raise ValueError("The model did not return valid JSON.")


def _to_score(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def normalize_result(data: dict) -> dict:
    """Make sure every field exists with the right type so the UI never crashes."""
    cats = data.get("category_scores") or {}
    category_scores = {k: _to_score(cats.get(k)) for k in CATEGORIES}

    overall = data.get("overall_score")
    if overall is None:
        overall = sum(category_scores.values()) / len(category_scores)

    jm = data.get("job_match_score")
    job_match = None if jm is None else _to_score(jm)

    def as_list(key):
        v = data.get(key)
        return v if isinstance(v, list) else []

    return {
        "overall_score": _to_score(overall),
        "summary": str(data.get("summary") or ""),
        "detected_role": str(data.get("detected_role") or "Not detected"),
        "category_scores": category_scores,
        "job_match_score": job_match,
        "strengths": [str(s) for s in as_list("strengths")],
        "issues": [i for i in as_list("issues") if isinstance(i, dict)],
        "improvements": [i for i in as_list("improvements") if isinstance(i, dict)],
        "missing_keywords": [str(k) for k in as_list("missing_keywords")],
        "rewritten_bullets": [b for b in as_list("rewritten_bullets") if isinstance(b, dict)],
    }


def analyze_resume(client, model: str, resume_text: str, job_description: str = "") -> dict:
    """Send the resume to Gemini and return a normalized result dict."""
    prompt = f"RESUME:\n\"\"\"\n{resume_text[:MAX_RESUME_CHARS]}\n\"\"\"\n"
    if job_description.strip():
        prompt += f"\nJOB DESCRIPTION:\n\"\"\"\n{job_description.strip()[:8000]}\n\"\"\"\n"
    else:
        prompt += "\nJOB DESCRIPTION: (none provided)\n"

    config = types.GenerateContentConfig(
        system_instruction=SYSTEM_PROMPT,
        response_mime_type="application/json",
    )
    candidates = [model] + [m for m in FALLBACK_MODELS if m != model]
    last_exc = None
    for name in candidates:
        try:
            response = client.models.generate_content(model=name, contents=prompt, config=config)
            break
        except Exception as exc:
            low = str(exc).lower()
            if "404" in low or "not found" in low or "no longer available" in low:
                last_exc = exc
                continue  # model retired/unknown -> try the next one
            raise
    else:
        raise last_exc
    return normalize_result(parse_json_response(response.text))


def friendly_error(exc: Exception) -> str:
    msg = str(exc)
    low = msg.lower()
    if "api key" in low or "api_key" in low or "permission" in low or "401" in low or "403" in low:
        return "Gemini rejected the API key. Please check that it is correct and active."
    if "429" in low or "quota" in low or "resource_exhausted" in low:
        return "Gemini rate limit / quota reached. Wait a minute and try again."
    if "404" in low or "not found" in low:
        return ("None of the Gemini models were found. Set the model in the sidebar to "
                "'gemini-flash-latest' or check https://ai.google.dev/gemini-api/docs/models")
    return f"Something went wrong while analyzing: {msg}"


# ----------------------------------------------------------------------------
# UI helpers
# ----------------------------------------------------------------------------
def score_color(score: int) -> str:
    if score >= 80:
        return "green"
    if score >= 60:
        return "orange"
    return "red"


def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent - likely to pass most ATS filters"
    if score >= 60:
        return "Good - but there is clear room to improve"
    if score >= 40:
        return "Fair - at risk of being filtered out"
    return "Poor - needs a major rework"


def get_api_key() -> str:
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass  # no secrets file locally
    return os.getenv("GEMINI_API_KEY", "")


SEVERITY_ICON = {"high": "🔴", "medium": "🟠", "low": "🟡"}


def render_results(result: dict) -> None:
    overall = result["overall_score"]

    st.divider()
    col1, col2 = st.columns([1, 2])
    with col1:
        st.markdown(
            f"<div style='text-align:center'>"
            f"<div style='font-size:64px;font-weight:700;color:{score_color(overall)};line-height:1'>"
            f"{overall}</div><div style='opacity:.7'>ATS score / 100</div></div>",
            unsafe_allow_html=True,
        )
    with col2:
        st.subheader(score_label(overall))
        st.write(result["summary"])
        st.caption(f"Detected target role: **{result['detected_role']}**")

    if result["job_match_score"] is not None:
        st.metric("Job description match", f"{result['job_match_score']}%")

    st.subheader("Score breakdown")
    for key, label in CATEGORIES.items():
        s = result["category_scores"][key]
        st.write(f"**{label}** - {s}/100")
        st.progress(s / 100)

    tabs = st.tabs(["Improvements", "Issues", "Strengths", "Keywords", "Rewritten bullets"])

    with tabs[0]:
        if not result["improvements"]:
            st.info("No improvements returned.")
        for imp in result["improvements"]:
            pr = str(imp.get("priority", "medium")).lower()
            st.markdown(f"{SEVERITY_ICON.get(pr, '🟠')} **{imp.get('suggestion', '')}**")
            if imp.get("example"):
                st.caption(f"Example: {imp['example']}")

    with tabs[1]:
        if not result["issues"]:
            st.success("No major issues found.")
        for iss in result["issues"]:
            sev = str(iss.get("severity", "medium")).lower()
            st.markdown(
                f"{SEVERITY_ICON.get(sev, '🟠')} **{iss.get('section', 'General')}** - "
                f"{iss.get('problem', '')}"
            )

    with tabs[2]:
        for s in result["strengths"] or ["No strengths returned."]:
            st.markdown(f"✅ {s}")

    with tabs[3]:
        if result["missing_keywords"]:
            st.write("Consider adding these keywords where they honestly apply to you:")
            st.write(" ".join(f"`{k}`" for k in result["missing_keywords"]))
        else:
            st.success("No missing keywords detected.")

    with tabs[4]:
        if not result["rewritten_bullets"]:
            st.info("No rewrites returned.")
        for b in result["rewritten_bullets"]:
            st.markdown(f"**Before:** {b.get('original', '')}")
            st.markdown(f"**After:** {b.get('improved', '')}")
            st.write("")

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


# ----------------------------------------------------------------------------
# Main app
# ----------------------------------------------------------------------------
def main() -> None:
    st.set_page_config(page_title="AI Resume ATS Analyzer", page_icon="📄", layout="centered")
    st.title("📄 AI Resume ATS Analyzer")
    st.write("Upload your resume to get an ATS score and actionable ways to improve it.")

    with st.sidebar:
        st.header("Settings")
        api_key = get_api_key()
        if api_key:
            st.success("API key loaded from secrets / environment.")
        else:
            api_key = st.text_input("Gemini API key", type="password",
                                    help="Get a free key at https://aistudio.google.com/apikey")
        model = st.text_input("Gemini model", value=DEFAULT_MODEL)
        st.caption("Your resume is sent to Google's Gemini API for analysis. "
                   "Don't upload anything you aren't comfortable sharing.")

    uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional)",
        height=150,
        placeholder="Paste the job posting here to get a job-match score and tailored keywords...",
    )

    if not st.button("Analyze resume", type="primary", disabled=uploaded is None):
        # show a previous result if the user interacts with other widgets
        if "result" in st.session_state:
            render_results(st.session_state["result"])
        return

    if not api_key:
        st.error("Please provide a Gemini API key in the sidebar.")
        return

    data = uploaded.getvalue()
    if len(data) > MAX_FILE_MB * 1024 * 1024:
        st.error(f"File is too large. Maximum size is {MAX_FILE_MB} MB.")
        return

    try:
        text = extract_text(uploaded.name, data)
    except Exception as exc:
        st.error(f"Could not read the file: {exc}")
        return

    if len(text) < MIN_RESUME_CHARS:
        st.error("Very little text could be extracted. If your resume is a scanned image, "
                 "that is also a problem for real ATS systems - export a text-based PDF or DOCX.")
        return

    cache_key = hashlib.sha256(f"{model}|{text}|{job_description}".encode()).hexdigest()
    cache = st.session_state.setdefault("cache", {})

    if cache_key in cache:
        result = cache[cache_key]
    else:
        try:
            client = genai.Client(api_key=api_key)
            with st.spinner("Analyzing your resume with Gemini..."):
                result = analyze_resume(client, model, text, job_description)
            cache[cache_key] = result
        except Exception as exc:
            st.error(friendly_error(exc))
            return

    st.session_state["result"] = result
    render_results(result)

    with st.expander("View extracted resume text (what an ATS would read)"):
        st.text(text[:5000] + ("..." if len(text) > 5000 else ""))


if __name__ == "__main__":
    main()
